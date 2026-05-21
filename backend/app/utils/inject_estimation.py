import openpyxl
import io
import os

TEMPLATE_MAP = {
    "CLOUD":     "Plan de Trabajo - Proyectos Cloud - PLANTILLA.xlsx",
    "DATA":      "Plan de Trabajo - Proyectos Data - PLANTILLA.xlsx",
    "SW_DEV":    "Plan de Trabajo - Proyectos Desarrollo de Software a la Medida - PLANTILLA.xlsx",
    "AI_AGENT":  "Plan de Trabajo - Proyectos Implementación de Agente de IA - PLANTILLA.xlsx",
    "ML":        "Plan de Trabajo - Proyectos Machine Learning Model(s) - PLANTILLA.xlsx",
}

TEMPLATES_DIR = os.path.join(os.path.dirname(__file__), "..", "..", "templates")


def _get_header_map(ws) -> dict:
    """Devuelve {nombre_columna: indice_1based} de la primera fila."""
    return {
        cell.value: cell.column
        for cell in ws[1]
        if cell.value is not None
    }


def _task_id_to_row(ws, id_col: int) -> dict:
    """Devuelve {task_id: row_number} para todas las filas de datos."""
    mapping = {}
    for row in ws.iter_rows(min_row=2):
        cell_val = row[id_col - 1].value
        if cell_val is not None:
            try:
                mapping[int(cell_val)] = row[0].row
            except (ValueError, TypeError):
                pass
    return mapping


def inject_estimation(estimation: dict) -> bytes:
    project_type = estimation.get("project_type", "").upper()
    template_file = TEMPLATE_MAP.get(project_type)
    if not template_file:
        raise ValueError(f"Tipo de proyecto desconocido: {project_type}")

    template_path = os.path.join(TEMPLATES_DIR, template_file)
    if not os.path.exists(template_path):
        raise FileNotFoundError(
            f"Plantilla no encontrada: {template_path}. "
            "Asegúrate de copiar los xlsx a backend/templates/"
        )

    wb = openpyxl.load_workbook(template_path)

    # ── Tabla_Tareas ──────────────────────────────────────────────────────────
    ws_tareas = wb["Tabla_Tareas"]
    h_t = _get_header_map(ws_tareas)
    nom_col  = h_t.get("Nombre", 4)
    dur_col  = h_t.get("Duración", 5)

    # Actualizar nombre del proyecto en fila 2 (ID 1)
    project_name = estimation.get("project_name", estimation.get("nombre_proyecto", ""))
    if project_name:
        ws_tareas.cell(row=2, column=nom_col).value = project_name

    # Índice nombre → fila para actualizar duraciones
    nombre_to_row = {}
    for row in ws_tareas.iter_rows(min_row=2):
        cell_val = row[nom_col - 1].value
        if cell_val:
            nombre_to_row[str(cell_val).lower()] = row[0].row

    # Calcular duración desde horas de asignaciones
    for asig in estimation.get("asignaciones", []):
        nombre = asig["nombre_tarea"].lower()
        horas = asig.get("horas", 0)
        if horas == 0:
            continue
        dias = max(1, round(horas / 8))
        # Match exacto primero, luego parcial
        row_num = nombre_to_row.get(nombre)
        if not row_num:
            for key, rn in nombre_to_row.items():
                if nombre in key or key in nombre:
                    row_num = rn
                    break
        if row_num:
            ws_tareas.cell(row=row_num, column=dur_col).value = f"{dias} días"

    # ── Tabla_Recursos ────────────────────────────────────────────────────────
    ws_rec = wb["Tabla_Recursos"]
    h_r = _get_header_map(ws_rec)
    nom_r_col    = h_r.get("Nombre", 2)
    tarifa_col   = h_r.get("Tasa estándar", 10)

    recursos_by_name = {
        r["nombre"].lower(): r
        for r in estimation.get("recursos", [])
    }

    for row in ws_rec.iter_rows(min_row=2):
        cell_val = row[nom_r_col - 1].value
        if not cell_val:
            continue
        cell_lower = str(cell_val).lower()
        for key, recurso in recursos_by_name.items():
            if key in cell_lower or cell_lower in key:
                tarifa = recurso.get("tarifa_hora", 0)
                row[tarifa_col - 1].value = f"${tarifa:.2f}/h"
                break

    # ── Tabla_asignación ──────────────────────────────────────────────────────
    ws_asig = wb["Tabla_asignación"]
    h_a = _get_header_map(ws_asig)
    tarea_col   = h_a.get("Nombre de tarea", 1)
    rec_col     = h_a.get("Nombre del recurso", 2)
    trabajo_col = h_a.get("Trabajo", 4)
    unid_col    = h_a.get("Unidades", 5)

    # Índice rápido de asignaciones: (tarea_lower, recurso_lower) → asignacion
    asig_index = {
        (a["nombre_tarea"].lower(), a["nombre_recurso"].lower()): a
        for a in estimation.get("asignaciones", [])
    }

    for row in ws_asig.iter_rows(min_row=2):
        tarea_val  = row[tarea_col - 1].value
        rec_val    = row[rec_col - 1].value
        if not tarea_val or not rec_val:
            continue

        tarea_lower = str(tarea_val).lower()
        rec_lower   = str(rec_val).lower()

        matched = None
        # Primero: match exacto
        matched = asig_index.get((tarea_lower, rec_lower))

        # Segundo: match parcial si no hubo exacto
        if not matched:
            for (at, ar), asig in asig_index.items():
                if (at in tarea_lower or tarea_lower in at) and \
                   (ar in rec_lower or rec_lower in ar):
                    matched = asig
                    break

        if matched:
            row[trabajo_col - 1].value  = f"{matched['horas']}h"
            row[unid_col - 1].value     = matched.get("unidades", "100%")

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()
